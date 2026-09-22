"""
tools/rag.py — RAG document teaching tool for NOVA.

Supports PDF, PPTX, TXT/MD.
Uses local ChromaDB + sentence-transformers embeddings.
Lazy-loads heavy deps so Streamlit startup stays fast.
"""

import logging
import os
import socket
import warnings
from typing import List, Dict, Any, Optional

# supress HF warning when loading. very cool
os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")
logging.getLogger("huggingface_hub").setLevel(logging.ERROR)
logging.getLogger("sentence_transformers").setLevel(logging.ERROR)
logging.getLogger("chromadb").setLevel(logging.ERROR)
warnings.filterwarnings("ignore", category=UserWarning, module="huggingface_hub")

socket.setdefaulttimeout(120.0)

CHROMA_DATA_PATH = os.environ.get("NOVA_CHROMA_PATH", "./chroma_db")
COLLECTION_NAME = "nova_documents"
_ACTIVE_DOC_PATH = os.path.join(CHROMA_DATA_PATH, "_active_doc.txt")

_client = None
_collection = None


def set_active_document(filename: str) -> None:
    """Mark this filename as the currently active document for ask_document()."""
    os.makedirs(CHROMA_DATA_PATH, exist_ok=True)
    with open(_ACTIVE_DOC_PATH, "w", encoding="utf-8") as f:
        f.write(filename.strip())


def get_active_document() -> str:
    """Return the filename of the currently active document, or '' if none."""
    try:
        with open(_ACTIVE_DOC_PATH, "r", encoding="utf-8") as f:
            return f.read().strip()
    except FileNotFoundError:
        return ""


def _get_embedding_function():
    """Load SentenceTransformer offline if cached, or download once cleanly."""
    from chromadb.utils import embedding_functions

    prev_offline = os.environ.get("HF_HUB_OFFLINE")
    try:
        os.environ["HF_HUB_OFFLINE"] = "1"
        return embedding_functions.SentenceTransformerEmbeddingFunction(
            model_name="all-MiniLM-L6-v2"
        )
    except Exception:
        # download nhi hua. allow download once
        if prev_offline is None:
            os.environ.pop("HF_HUB_OFFLINE", None)
        else:
            os.environ["HF_HUB_OFFLINE"] = prev_offline

        try:
            return embedding_functions.SentenceTransformerEmbeddingFunction(
                model_name="all-MiniLM-L6-v2"
            )
        except Exception:
            return embedding_functions.DefaultEmbeddingFunction()
    finally:
        if prev_offline is not None:
            os.environ["HF_HUB_OFFLINE"] = prev_offline


def _get_collection():
    """Lazy init Chroma collection; auto-heal embedding conflicts."""
    global _client, _collection
    if _collection is not None:
        return _collection

    import chromadb

    os.makedirs(CHROMA_DATA_PATH, exist_ok=True)
    _client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)

    embedding_fn = _get_embedding_function()

    try:
        _collection = _client.get_or_create_collection(
            name=COLLECTION_NAME,
            embedding_function=embedding_fn,
        )
    except Exception:
        # Old DB used a different embedding function — reset collection once.
        try:
            _client.delete_collection(name=COLLECTION_NAME)
        except Exception:
            pass
        _collection = _client.get_or_create_collection(
            name=COLLECTION_NAME,
            embedding_function=embedding_fn,
        )

    return _collection


def _extract_sections(file_path: str) -> List[Dict[str, str]]:
    """Extract raw text from PDF, PPTX, or TXT/MD."""
    ext = os.path.splitext(file_path)[1].lower()
    sections = []

    try:
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides, start=1):
                slide_content = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text_val = shape.text.strip()
                        if text_val:
                            slide_content.append(text_val)
                    if hasattr(shape, "has_table") and shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                            if row_text:
                                slide_content.append(row_text)
                if hasattr(slide, "notes_slide") and slide.notes_slide:
                    notes_tf = slide.notes_slide.notes_text_frame
                    if notes_tf and notes_tf.text.strip():
                        slide_content.append(f"[Notes: {notes_tf.text.strip()}]")

                full_slide_text = "\n".join(slide_content).strip()
                if full_slide_text:
                    sections.append({"text": full_slide_text, "source_label": f"Slide {i}"})
            return sections

        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            for i, page in enumerate(reader.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    sections.append({"text": page_text, "source_label": f"Page {i}"})
            return sections

        if ext in (".txt", ".md"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
            if content:
                # split large plain text files by double newlines
                paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
                if not paragraphs:
                    paragraphs = [content]
                for idx, para in enumerate(paragraphs, start=1):
                    sections.append({"text": para, "source_label": f"Section {idx}"})
            return sections

        return [{"text": f"Unsupported file type '{ext}'. Supported: .pptx, .pdf, .txt, .md", "source_label": "Error"}]

    except Exception as error:
        return [{"text": f"Error reading file '{file_path}': {error}", "source_label": "Error"}]


def _chunk_sections(sections: List[Dict[str, str]], chunk_size: int = 1000, overlap: int = 150) -> List[Dict[str, str]]:
    """Split text into overlapping chunks."""
    chunks = []
    step = max(100, chunk_size - overlap)

    for section in sections:
        text = section["text"]
        label = section["source_label"]

        if len(text) <= chunk_size:
            chunks.append({"text": text, "source_label": label})
            continue

        start = 0
        text_len = len(text)
        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk_slice = text[start:end].strip()
            if chunk_slice:
                chunks.append({"text": chunk_slice, "source_label": label})
            if end >= text_len:
                break
            start += step

    return chunks


def ingest_document(file_path: str, semester: str = "", subject: str = "") -> str:
    """Read a document, chunk it, embed it, and store it in ChromaDB.

    Args:
        file_path: Local path to the uploaded document file.
        semester: Optional semester tag (e.g. '3').
        subject: Optional subject tag (e.g. 'Theory of Computation').
    """
    if not file_path or not str(file_path).strip():
        return "File path cannot be empty."

    if not os.path.exists(file_path):
        return f"File does not exist: {file_path}"

    filename = os.path.basename(file_path)
    sections = _extract_sections(file_path)

    if not sections:
        return f"No extractable text found in '{filename}'."

    first_text = sections[0]["text"]
    if first_text.startswith("Unsupported") or first_text.startswith("Error"):
        return first_text

    chunks = _chunk_sections(sections)
    if not chunks:
        return f"Could not create text chunks from '{filename}'."

    ids = [f"{filename}_chunk_{i}" for i in range(len(chunks))]
    documents = [c["text"] for c in chunks]
    metadatas = [
        {
            "filename": filename,
            "source_label": c["source_label"],
            "semester": str(semester or "").strip(),
            "subject": str(subject or "").strip(),
            "chunk_index": i,
        }
        for i, c in enumerate(chunks)
    ]

    try:
        collection = _get_collection()
        try:
            collection.delete(where={"filename": filename})
        except Exception:
            pass

        collection.add(documents=documents, ids=ids, metadatas=metadatas)
        set_active_document(filename)
        return (
            f"Successfully processed '{filename}' into {len(chunks)} chunks "
            "in vector store."
        )
    except Exception as error:
        return f"Failed to store embeddings for '{filename}': {error}"


def get_indexed_documents() -> List[Dict[str, Any]]:
    """return a list of unique indexed documents and their chunk counts"""
    try:
        collection = _get_collection()
        data = collection.get(include=["metadatas"])
        metadatas = data.get("metadatas") or []
        doc_stats: Dict[str, Dict[str, Any]] = {}
        for meta in metadatas:
            if not meta or "filename" not in meta:
                continue
            fname = meta["filename"]
            if fname not in doc_stats:
                doc_stats[fname] = {
                    "filename": fname,
                    "chunks": 0,
                    "semester": meta.get("semester", ""),
                    "subject": meta.get("subject", ""),
                }
            doc_stats[fname]["chunks"] += 1
        return list(doc_stats.values())
    except Exception:
        return []


def delete_document(filename: str) -> bool:
    """delete a document and its chunks from chroma"""
    try:
        collection = _get_collection()
        collection.delete(where={"filename": filename})
        active = get_active_document()
        if active == filename:
            docs = get_indexed_documents()
            new_active = docs[0]["filename"] if docs else ""
            set_active_document(new_active)
        return True
    except Exception:
        return False


def clear_all_documents() -> bool:
    """clear all documents from chroma"""
    global _collection
    try:
        collection = _get_collection()
        _client.delete_collection(name=COLLECTION_NAME)
        _collection = None
        set_active_document("")
        return True
    except Exception:
        return False


def ask_document(question: str, filename: str = "") -> str:
    """Search uploaded study materials and return relevant context.

    Args:
        question: The user's query about the uploaded document.
        filename: Optional specific document filename to search within.
    """
    if not question or not str(question).strip():
        return "Please ask a specific question about the document."

    try:
        collection = _get_collection()
        total_chunks = collection.count()
        if total_chunks == 0:
            return (
                "FALLBACK:NO_DOCS | No study materials have been uploaded yet. "
                "Please upload a PPTX, PDF, TXT, or MD file using the sidebar first."
            )

        # determine target file
        target_file = filename.strip() if filename else get_active_document()
        if target_file in ("", "All Documents", "all"):
            target_file = None

        query_kwargs: Dict[str, Any] = {
            "query_texts": [question],
            "n_results": min(4, total_chunks),
        }
        if target_file:
            query_kwargs["where"] = {"filename": target_file}

        try:
            results = collection.query(**query_kwargs)
        except Exception:
            # if filtered query failed (filename not in collection), retry across all docs
            query_kwargs.pop("where", None)
            results = collection.query(**query_kwargs)

        docs = (results.get("documents") or [[]])[0]
        metas = (results.get("metadatas") or [[]])[0]

        if not docs:
            return (
                "FALLBACK:EMPTY_RETRIEVAL | I searched the uploaded materials "
                f"for '{question}' but found no strongly matching content. "
                "Please try rephrasing or checking another uploaded document."
            )

        pieces = []
        for doc, meta in zip(docs, metas or [{}] * len(docs)):
            src = (meta or {}).get("filename", "Document")
            label = (meta or {}).get("source_label", "")
            header = f"[{src} - {label}]" if label else f"[{src}]"
            pieces.append(f"{header}:\n{doc}")

        return "Retrieved Context from Uploaded Files:\n\n" + "\n\n---\n\n".join(pieces)

    except Exception as error:
        return (
            "FALLBACK:RAG_ERROR | Document search failed "
            f"({type(error).__name__}: {error}). "
            "I can still answer generally without the file."
        )